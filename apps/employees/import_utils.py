"""Import en masse d'employés — CDC §11.5. CSV et Excel (.xlsx) supportés ;
le format .xls (binaire, pré-2007) n'est pas pris en charge — openpyxl ne le
lit pas et ajouter xlrd pour ce seul format legacy n'est pas justifié.
"""

import csv
import io
import secrets
from datetime import date as date_cls, datetime, timedelta
from decimal import Decimal, InvalidOperation

from django.contrib.auth import get_user_model
from django.core.mail import send_mail
from django.db import transaction
from django.utils import timezone

from apps.accounts.password_policy import record_password
from apps.agencies.models import Agency
from apps.departments.models import Department, Position

from .forms import CREATABLE_ROLES, MANAGER_ROLES
from .models import Employee

User = get_user_model()

COLUMNS = [
    "E-mail", "Prénom", "Nom", "Rôle", "Département", "Poste", "Agence principale",
    "Manager (e-mail)", "Type de contrat", "Date d'entrée", "Droit aux congés (j/an)", "Matricule",
]

TEMPLATE_EXAMPLE_ROW = [
    "jean.dupont@example.com", "Jean", "Dupont", "Employé", "Ventes", "Vendeur",
    "Siège Social", "", "CDI", "15/01/2026", "25", "",
]

ROLE_LOOKUP = {label.upper(): code for code, label in CREATABLE_ROLES}
ROLE_LOOKUP.update({code: code for code, _ in CREATABLE_ROLES})
CONTRACT_LOOKUP = {label.upper(): code for code, label in Employee.ContractType.choices}
CONTRACT_LOOKUP.update({code: code for code, _ in Employee.ContractType.choices})

DATE_FORMATS = ["%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"]


class ImportFileError(Exception):
    pass


def parse_upload(uploaded_file):
    """Retourne une liste de dicts {en-tête: valeur} — une entrée par ligne
    de données (hors en-tête)."""
    name = uploaded_file.name.lower()
    if name.endswith(".csv"):
        decoded = uploaded_file.read().decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(decoded))
        return [dict(row) for row in reader]
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        wb = load_workbook(uploaded_file, data_only=True)
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        try:
            headers = [str(h).strip() if h is not None else "" for h in next(rows_iter)]
        except StopIteration:
            return []
        result = []
        for values in rows_iter:
            if all(v is None or str(v).strip() == "" for v in values):
                continue
            result.append({headers[i]: values[i] for i in range(len(headers)) if i < len(values)})
        return result
    raise ImportFileError("Format non supporté, utilisez un fichier .csv ou .xlsx.")


def _parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date_cls):
        return value
    value = str(value).strip()
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            continue
    return None


def _clean_str(value):
    if value is None:
        return ""
    return str(value).strip()


def validate_row(raw, tenant, row_number, seen_emails, seen_matricules):
    """Valide une ligne brute. Retourne (donnée_nettoyée_ou_None, [erreurs])."""
    errors = []

    email = _clean_str(raw.get("E-mail")).lower()
    first_name = _clean_str(raw.get("Prénom"))
    last_name = _clean_str(raw.get("Nom"))
    role_raw = _clean_str(raw.get("Rôle"))
    dept_raw = _clean_str(raw.get("Département"))
    pos_raw = _clean_str(raw.get("Poste"))
    agency_raw = _clean_str(raw.get("Agence principale"))
    manager_raw = _clean_str(raw.get("Manager (e-mail)")).lower()
    contract_raw = _clean_str(raw.get("Type de contrat"))
    leave_days_raw = _clean_str(raw.get("Droit aux congés (j/an)"))
    matricule_raw = _clean_str(raw.get("Matricule"))

    if not email:
        errors.append("E-mail manquant.")
    elif "@" not in email:
        errors.append("E-mail invalide.")
    elif email in seen_emails:
        errors.append("E-mail en double dans le fichier.")
    elif User.objects.filter(email=email).exists():
        errors.append("Un compte existe déjà avec cet e-mail.")

    if not first_name:
        errors.append("Prénom manquant.")
    if not last_name:
        errors.append("Nom manquant.")

    role = ROLE_LOOKUP.get(role_raw.upper())
    if not role:
        errors.append(f"Rôle invalide : « {role_raw} » (Employé, Manager ou Administrateur).")

    department = None
    if dept_raw:
        department = Department.objects.all_tenants().filter(
            tenant=tenant, name__iexact=dept_raw, is_active=True
        ).first()
        if department is None:
            errors.append(f"Département introuvable : « {dept_raw} ».")

    position = None
    if pos_raw:
        position = Position.objects.all_tenants().filter(
            tenant=tenant, title__iexact=pos_raw, is_active=True
        ).first()
        if position is None:
            errors.append(f"Poste introuvable : « {pos_raw} ».")

    agency = None
    if not agency_raw:
        errors.append("Agence principale manquante.")
    else:
        agency = Agency.objects.all_tenants().filter(tenant=tenant, name__iexact=agency_raw, is_active=True).first()
        if agency is None:
            errors.append(f"Agence introuvable : « {agency_raw} ».")

    manager = None
    if manager_raw:
        manager = User.objects.filter(tenant=tenant, email__iexact=manager_raw, role__in=MANAGER_ROLES).first()
        if manager is None:
            errors.append(f"Manager introuvable (ou rôle non habilité) : « {manager_raw} ».")

    contract_type = CONTRACT_LOOKUP.get(contract_raw.upper())
    if not contract_type:
        errors.append(f"Type de contrat invalide : « {contract_raw} ».")

    hire_date_val = raw.get("Date d'entrée")
    hire_date = _parse_date(hire_date_val) if hire_date_val not in (None, "") else None
    if hire_date is None:
        errors.append(f"Date d'entrée invalide ou manquante : « {_clean_str(hire_date_val)} » (JJ/MM/AAAA).")
    elif hire_date > timezone.localdate() + timedelta(days=30):
        errors.append("La date d'entrée ne peut pas être à plus de 30 jours dans le futur.")

    leave_days = Decimal("25")
    if leave_days_raw:
        try:
            leave_days = Decimal(leave_days_raw.replace(",", "."))
        except InvalidOperation:
            errors.append(f"Droit aux congés invalide : « {leave_days_raw} ».")

    if matricule_raw:
        if matricule_raw in seen_matricules:
            errors.append("Matricule en double dans le fichier.")
        elif Employee.objects.all_tenants().filter(tenant=tenant, matricule=matricule_raw).exists():
            errors.append(f"Matricule déjà utilisé : « {matricule_raw} ».")

    if errors:
        return None, errors

    seen_emails.add(email)
    if matricule_raw:
        seen_matricules.add(matricule_raw)
    return {
        "row": row_number,
        "email": email,
        "first_name": first_name,
        "last_name": last_name,
        "role": role,
        "department_id": str(department.id) if department else "",
        "department_name": department.name if department else "",
        "position_id": str(position.id) if position else "",
        "position_name": position.title if position else "",
        "agency_id": str(agency.id),
        "agency_name": agency.name,
        "manager_id": str(manager.id) if manager else "",
        "contract_type": contract_type,
        "hire_date": hire_date.isoformat(),
        "annual_leave_days": str(leave_days),
        "matricule": matricule_raw,
    }, []


def validate_rows(raw_rows, tenant):
    """Valide toutes les lignes. Retourne (lignes_valides, lignes_en_erreur)."""
    valid, invalid = [], []
    seen_emails = set()
    seen_matricules = set()
    for i, raw in enumerate(raw_rows, start=2):  # ligne 1 = en-tête
        cleaned, errors = validate_row(raw, tenant, i, seen_emails, seen_matricules)
        if errors:
            invalid.append({"row": i, "email": _clean_str(raw.get("E-mail")) or "—", "errors": errors})
        else:
            valid.append(cleaned)
    return valid, invalid


def commit_import(tenant, validated_rows):
    """CDC §11.3.2 : crée chaque compte + profil, envoie l'e-mail de bienvenue.
    Une ligne qui échoue à la création (cas limite non détecté en validation)
    n'annule pas les autres — chacune est isolée dans sa propre transaction."""
    created, failed = [], []
    for data in validated_rows:
        try:
            with transaction.atomic():
                temp_password = secrets.token_urlsafe(10)
                user = User.objects.create_user(
                    email=data["email"],
                    password=temp_password,
                    tenant=tenant,
                    role=data["role"],
                    first_name=data["first_name"],
                    last_name=data["last_name"],
                    must_change_password=True,
                )
                record_password(user)
                employee = Employee(
                    tenant=tenant,
                    user=user,
                    department_id=data["department_id"] or None,
                    position_id=data["position_id"] or None,
                    primary_agency_id=data["agency_id"],
                    manager_id=data["manager_id"] or None,
                    contract_type=data["contract_type"],
                    hire_date=date_cls.fromisoformat(data["hire_date"]),
                    annual_leave_days=Decimal(data["annual_leave_days"]),
                    matricule=data["matricule"],
                )
                employee.full_clean()
                employee.save()
        except Exception as exc:  # noqa: BLE001 — isolation voulue, cf. docstring
            failed.append({"row": data["row"], "email": data["email"], "error": str(exc)})
            continue

        email_sent = True
        try:
            # Le compte est déjà committé à ce stade : un échec d'envoi (SMTP
            # indisponible, etc.) ne doit ni annuler la ligne (elle a bien été
            # créée) ni interrompre le traitement des lignes suivantes — sans
            # ce try/except, une exception ici remontait hors de la boucle et
            # cassait l'isolation par ligne promise ci-dessus.
            send_mail(
                subject="Bienvenue sur GeoPresence",
                message=(
                    f"Bonjour {user.first_name},\n\n"
                    f"Votre compte a été créé sur GeoPresence.\n"
                    f"E-mail : {user.email}\n"
                    f"Mot de passe temporaire : {temp_password}\n\n"
                    "Connectez-vous et changez votre mot de passe dès la première connexion."
                ),
                from_email=None,
                recipient_list=[user.email],
            )
        except Exception:  # noqa: BLE001
            email_sent = False
        created.append({"row": data["row"], "email": data["email"], "email_sent": email_sent})

    return created, failed
